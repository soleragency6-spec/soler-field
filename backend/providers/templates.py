"""Uploaded client template analysis (reports / invoices / quotes).

Extracts structure from a professional's existing PDF/DOCX so future documents
reproduce their layout & section order. Falls back to a deterministic heuristic
when no AI provider is configured (the `ai.analyze_template` hook can later
enrich this via OpenAI/vision without changing callers).
"""
import io
import re


def _clean(line: str) -> str:
    return re.sub(r"\s+", " ", line or "").strip()


def _looks_like_title(line: str) -> bool:
    l = _clean(line)
    if not l or len(l) > 60:
        return False
    letters = [c for c in l if c.isalpha()]
    if len(letters) < 2:
        return False
    upper_ratio = sum(1 for c in letters if c.isupper()) / len(letters)
    ends_colon = l.endswith(":")
    titlecase = l[0].isupper() and len(l.split()) <= 6
    return upper_ratio > 0.6 or ends_colon or titlecase


def analyze(data: bytes, mime: str, filename: str) -> dict:
    name = (filename or "").lower()
    kind_src = "image"
    text = ""
    tables = 0
    try:
        if name.endswith(".pdf") or "pdf" in (mime or ""):
            kind_src = "pdf"
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages = []
            for pg in reader.pages[:5]:
                pages.append(pg.extract_text() or "")
            text = "\n".join(pages)
        elif name.endswith(".docx") or "word" in (mime or "") or "officedocument" in (mime or ""):
            kind_src = "docx"
            from docx import Document

            doc = Document(io.BytesIO(data))
            lines = []
            titles_from_headings = []
            for p in doc.paragraphs:
                t = _clean(p.text)
                if not t:
                    continue
                lines.append(t)
                if (p.style and p.style.name or "").lower().startswith("heading"):
                    titles_from_headings.append(t)
            tables = len(doc.tables)
            text = "\n".join(lines)
            if titles_from_headings:
                return _build(kind_src, titles_from_headings, text, tables)
    except Exception:  # noqa: BLE001
        text = ""

    titles = []
    for line in text.splitlines():
        if _looks_like_title(line):
            titles.append(_clean(line).rstrip(":"))
    # de-dup keep order
    seen = set()
    titles = [t for t in titles if not (t in seen or seen.add(t))][:12]
    return _build(kind_src, titles, text, tables)


def _build(source: str, titles: list[str], text: str, tables: int) -> dict:
    sections = [{"title": t, "body": ""} for t in titles]
    return {
        "source": source,
        "sections": sections,
        "detected": {"section_count": len(sections), "tables": tables},
        "excerpt": (text or "").strip()[:1500],
    }
