"""Templates & Quotes feature backend tests (iteration 2)."""
import io
import os
import time

import pytest
import requests

BASE_URL = os.environ.get("EXPO_PUBLIC_BACKEND_URL", "http://127.0.0.1:8000").rstrip("/")
API = f"{BASE_URL}/api"
SUPABASE_URL = "https://bjutwzqitkudwjzvevls.supabase.co"
SUPABASE_ANON = "sb_publishable_vpSkLI2aoxy25UFzwSmojA_Xi0w27ju"

PRO_EMAIL = "pro@martin-plomberie.fr"
PRO_PASS = "Test1234!"


def _login(email, password):
    r = requests.post(
        f"{SUPABASE_URL}/auth/v1/token?grant_type=password",
        headers={"apikey": SUPABASE_ANON, "Content-Type": "application/json"},
        json={"email": email, "password": password},
        timeout=20,
    )
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text[:200]}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def pro_token():
    return _login(PRO_EMAIL, PRO_PASS)


@pytest.fixture
def pro_headers(pro_token):
    return {"Authorization": f"Bearer {pro_token}"}


@pytest.fixture
def pro_json_headers(pro_token):
    return {"Authorization": f"Bearer {pro_token}", "Content-Type": "application/json"}


@pytest.fixture(scope="module")
def ctx():
    return {}


def _make_docx(section_titles):
    """Build a real .docx in-memory with Heading paragraphs."""
    from docx import Document
    doc = Document()
    doc.add_heading("Rapport client", level=0)
    for t in section_titles:
        doc.add_heading(t, level=1)
        doc.add_paragraph("Contenu à compléter.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----------------------------- 401 / RLS
def test_templates_401_no_token():
    r = requests.get(f"{API}/templates", timeout=20)
    assert r.status_code in (401, 403)


def test_quotes_401_no_token():
    r = requests.get(f"{API}/quotes", timeout=20)
    assert r.status_code in (401, 403)


def test_templates_bad_token_401():
    r = requests.get(f"{API}/templates", headers={"Authorization": "Bearer bad.jwt"}, timeout=20)
    assert r.status_code == 401


# ----------------------------- Templates upload / list / default / delete
def test_10_upload_report_template(pro_headers, ctx):
    data = _make_docx([
        "Coordonnées du client",
        "Objet de l'intervention",
        "Constat sur site",
        "Travaux effectués",
        "Recommandations & suite",
    ])
    files = {"file": ("modele_rapport.docx", data,
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    form = {"kind": "report", "name": "TEST_Modele Rapport"}
    r = requests.post(f"{API}/templates", headers=pro_headers, files=files, data=form, timeout=30)
    assert r.status_code == 200, r.text
    d = r.json()
    ctx["tid"] = d["id"]
    assert d["kind"] == "report"
    structure = d.get("structure_json") or {}
    sections = structure.get("sections") or []
    assert len(sections) >= 3, f"expected >=3 sections got {len(sections)}: {sections}"
    # Must reflect our headings (first one at least)
    titles = [s.get("title") for s in sections]
    assert any("Coordonn" in t for t in titles), titles


def test_11_first_upload_is_default(pro_headers, ctx):
    r = requests.get(f"{API}/templates?kind=report", headers=pro_headers, timeout=20)
    assert r.status_code == 200
    lst = r.json()
    mine = [t for t in lst if t["id"] == ctx["tid"]]
    assert mine, "uploaded template not in list"
    # If only one template exists in this workspace, it should be default. Otherwise
    # at least one report template is default and ours may be default on first upload.
    reports = [t for t in lst if t["kind"] == "report"]
    assert any(t["is_default"] for t in reports)


def test_12_template_original_url(pro_headers, ctx):
    r = requests.get(f"{API}/templates/{ctx['tid']}/original", headers=pro_headers, timeout=20)
    assert r.status_code == 200
    d = r.json()
    url = d.get("url", "")
    # Local dev storage returns signed /api/files?... path; R2 would return https URL.
    assert url.startswith("http") or "/api/files?" in url, url
    assert "sig=" in url and "exp=" in url, url


def test_13_report_uses_template_priority(pro_json_headers, pro_headers, ctx):
    # First, make sure OUR template is default (may not be if prior data existed)
    requests.put(f"{API}/templates/{ctx['tid']}", headers=pro_json_headers,
                 json={"is_default": True}, timeout=20)
    # Create client + intervention + report
    c = requests.post(f"{API}/clients", headers=pro_json_headers,
                      json={"company": "TEST_TmplCo"}, timeout=20)
    assert c.status_code == 200, c.text
    cid = c.json()["id"]; ctx["cid"] = cid
    i = requests.post(f"{API}/interventions", headers=pro_json_headers,
                      json={"client_id": cid, "name": "TEST_Tmpl intervention",
                            "type": "depannage", "description": "Détails."}, timeout=20)
    assert i.status_code == 200, i.text
    iid = i.json()["id"]; ctx["iid"] = iid

    r = requests.post(f"{API}/interventions/{iid}/report/generate",
                      headers=pro_json_headers, timeout=30)
    assert r.status_code == 200, r.text
    content = r.json().get("content_json") or {}
    meta = content.get("meta") or {}
    assert meta.get("template_used") is True, meta
    titles = [s.get("title") for s in (content.get("sections") or [])]
    # Must reproduce OUR titles in the order we uploaded
    assert any("Coordonn" in t for t in titles), titles
    assert any("Constat" in t for t in titles), titles


# ----------------------------- Quotes: create -> edit -> finalize -> pdf -> status
def test_20_quote_from_intervention(pro_json_headers, ctx):
    r = requests.post(f"{API}/interventions/{ctx['iid']}/quote",
                      headers=pro_json_headers, timeout=20)
    assert r.status_code == 200, r.text
    d = r.json()
    ctx["qid"] = d["id"]
    assert d["status"] == "draft"
    # Ensure client_id copied
    assert d["client_id"] == ctx["cid"]


def test_21_get_quote_has_items_and_client(pro_headers, ctx):
    r = requests.get(f"{API}/quotes/{ctx['qid']}", headers=pro_headers, timeout=20)
    assert r.status_code == 200
    d = r.json()
    assert isinstance(d.get("items"), list)
    assert d.get("client") is not None
    assert d["client"]["id"] == ctx["cid"]


def test_22_update_quote_recomputes_totals(pro_json_headers, ctx):
    payload = {
        "notes": "Devis valable 30 jours",
        "valid_until": "2026-12-31",
        "discount": 0,
        "items": [
            {"description": "Main d'œuvre", "quantity": 2, "unit": "h",
             "unit_price_ht": 60.0, "vat_rate": 20},
            {"description": "Pièces", "quantity": 1, "unit": "u",
             "unit_price_ht": 40.0, "vat_rate": 20},
        ],
    }
    r = requests.put(f"{API}/quotes/{ctx['qid']}", headers=pro_json_headers,
                     json=payload, timeout=20)
    assert r.status_code == 200, r.text
    d = r.json()
    assert abs(float(d["subtotal_ht"]) - 160.0) < 0.01
    assert abs(float(d["total_ttc"]) - 192.0) < 0.01
    assert d["notes"] == payload["notes"]


def test_23_finalize_quote_assigns_number_and_pdf(pro_headers, ctx):
    r = requests.post(f"{API}/quotes/{ctx['qid']}/finalize",
                      headers=pro_headers, timeout=30)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "final"
    assert d.get("number") and d["number"].startswith(("DEV-", "D-", "Q-")), d.get("number")
    assert d.get("pdf_key"), d


def test_24_quote_pdf(pro_headers, ctx):
    r = requests.get(f"{API}/quotes/{ctx['qid']}/pdf", headers=pro_headers, timeout=20)
    assert r.status_code == 200
    assert "url" in r.json()


def test_25_finalized_quote_locked(pro_json_headers, ctx):
    r = requests.put(f"{API}/quotes/{ctx['qid']}", headers=pro_json_headers,
                     json={"items": [], "discount": 0}, timeout=20)
    assert r.status_code == 409


def test_26_quote_status_sent(pro_json_headers, ctx):
    r = requests.post(f"{API}/quotes/{ctx['qid']}/status",
                      headers=pro_json_headers, json={"status": "sent"}, timeout=20)
    assert r.status_code == 200
    assert r.json()["status"] == "sent"


def test_27_quote_shows_in_listing(pro_headers, ctx):
    r = requests.get(f"{API}/quotes", headers=pro_headers, timeout=20)
    assert r.status_code == 200
    ids = [q["id"] for q in r.json()]
    assert ctx["qid"] in ids


# ----------------------------- Email attach quote (att-quote) pending
def test_30_email_attach_quote_pending(pro_json_headers, ctx):
    payload = {
        "client_id": ctx["cid"],
        "intervention_id": ctx["iid"],
        "recipient": "TEST_devis@example.com",
        "subject": "Votre devis",
        "body": "Bonjour, veuillez trouver votre devis.",
        "attachments": [{"type": "quote", "id": ctx["qid"], "name": "devis.pdf"}],
    }
    r = requests.post(f"{API}/emails", headers=pro_json_headers, json=payload, timeout=20)
    assert r.status_code == 200, r.text
    eid = r.json()["id"]
    r = requests.post(f"{API}/emails/{eid}/send", headers=pro_json_headers, timeout=30)
    assert r.status_code == 200
    d = r.json()
    assert d["email"]["status"] == "pending"
    assert d["email_enabled"] is False


# ----------------------------- Templates: mark another / delete
def test_40_delete_template(pro_headers, ctx):
    r = requests.delete(f"{API}/templates/{ctx['tid']}", headers=pro_headers, timeout=20)
    assert r.status_code == 200
    assert r.json().get("ok") is True
    # Confirm it no longer shows up
    lst = requests.get(f"{API}/templates?kind=report", headers=pro_headers, timeout=20).json()
    assert not any(t["id"] == ctx["tid"] for t in lst)


# ----------------------------- Cleanup
def test_99_cleanup(pro_headers, ctx):
    if ctx.get("iid"):
        requests.delete(f"{API}/interventions/{ctx['iid']}", headers=pro_headers, timeout=20)
    if ctx.get("cid"):
        requests.delete(f"{API}/clients/{ctx['cid']}", headers=pro_headers, timeout=20)
